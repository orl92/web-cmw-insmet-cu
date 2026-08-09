import datetime
from io import BytesIO

from django.test import SimpleTestCase
from openpyxl import load_workbook

from apps.meteo.utils import excel_forecast as ef


class ExcelForecastUtilsTest(SimpleTestCase):
    def test_to_date_accepts_dd_mm_yyyy(self):
        self.assertEqual(ef._to_date('09/08/2026'), datetime.date(2026, 8, 9))

    def test_to_date_accepts_iso(self):
        self.assertEqual(ef._to_date('2026-08-09'), datetime.date(2026, 8, 9))

    def test_to_date_accepts_datetime_and_date(self):
        self.assertEqual(
            ef._to_date(datetime.datetime(2026, 8, 9, 12, 0)), datetime.date(2026, 8, 9)
        )
        self.assertEqual(ef._to_date(datetime.date(2026, 8, 9)), datetime.date(2026, 8, 9))

    def test_to_date_invalid_returns_none(self):
        self.assertIsNone(ef._to_date('no-es-fecha'))
        self.assertIsNone(ef._to_date(None))

    def test_to_time_serial_float(self):
        self.assertEqual(ef._to_time(0.3125), '07:30')
        self.assertEqual(ef._to_time(0.5), '12:00')

    def test_to_time_serial_int(self):
        self.assertEqual(ef._to_time(0), '00:00')

    def test_to_time_out_of_range_serial_returns_empty(self):
        self.assertEqual(ef._to_time(1.5), '')
        self.assertEqual(ef._to_time(-0.1), '')

    def test_to_time_time_objects(self):
        self.assertEqual(ef._to_time(datetime.time(7, 30)), '07:30')
        self.assertEqual(ef._to_time(datetime.datetime(2026, 8, 9, 14, 30)), '14:30')

    def test_to_time_strings(self):
        self.assertEqual(ef._to_time('14:30'), '14:30')
        self.assertEqual(ef._to_time('14:30:00'), '14:30')
        self.assertEqual(ef._to_time('auto'), '')
        self.assertEqual(ef._to_time(''), '')


class BuildTemplateTest(SimpleTestCase):
    def test_group_headers_survive_merge(self):
        wb = load_workbook(BytesIO(ef.build_template()))
        ws = wb.active
        expected = {
            'B2': 'Temperatura',
            'E2': 'Tiempo',
            'H2': 'Viento (dd)',
            'K2': 'Viento (ff)',
            'N2': 'Mar',
        }
        for cell, header in expected.items():
            self.assertEqual(ws[cell].value, header)

    def test_group_headers_merged_in_threes(self):
        wb = load_workbook(BytesIO(ef.build_template()))
        ws = wb.active
        merged = {str(r) for r in ws.merged_cells.ranges}
        self.assertIn('B2:D2', merged)
        self.assertIn('E2:G2', merged)
        self.assertIn('H2:J2', merged)
        self.assertIn('K2:M2', merged)
        self.assertIn('N2:P2', merged)
