import datetime

import openpyxl
from django.contrib.auth.mixins import LoginRequiredMixin, PermissionRequiredMixin
from django.http import HttpResponse
from django.utils import timezone
from django.views import View
from openpyxl.styles import Alignment, Border, Font, Side

from common.views import CSVExportView
from dashboard.models import (
    Customer,
    EarlyWarning,
    Forecasts,
    Invoice,
    Service,
    ServiceSubscription,
    StormWarning,
    TropicalCyclone,
    WeatherReport,
)


class CustomerCSVExportView(CSVExportView):
    model = Customer
    permission_required = 'dashboard.view_customer'
    filename = 'clientes.csv'
    columns = [
        ('Empresa', 'company_name'),
        ('REEUP', 'reeup'),
        ('NIT', 'nit'),
        ('Cuenta Bancaria', 'account'),
        ('Teléfono', 'phone'),
        ('Dirección', 'address'),
        ('Email', lambda o: o.user.email if o.user else ''),
        ('Registro Activo', lambda o: 'Sí' if o.record_active else 'No'),
    ]


class ServiceCSVExportView(CSVExportView):
    model = Service
    permission_required = 'dashboard.view_service'
    filename = 'servicios.csv'
    columns = [
        ('Título', 'title'),
        ('Precio', 'price'),
        ('Tipo', lambda o: o.get_service_type_display()),
        ('Resumen', 'summary'),
        ('Registro Activo', lambda o: 'Sí' if o.record_active else 'No'),
    ]


class ServiceSubscriptionCSVExportView(CSVExportView):
    model = ServiceSubscription
    permission_required = 'dashboard.view_subscription'
    filename = 'suscripciones.csv'
    columns = [
        ('Cliente', lambda o: o.customer.company_name if o.customer else ''),
        ('Servicio', lambda o: o.service.title if o.service else ''),
        ('Inicio', lambda o: o.start_date.isoformat() if o.start_date else ''),
        ('Fin', lambda o: o.end_date.isoformat() if o.end_date else ''),
        ('Estado Pago', lambda o: o.get_payment_status_display()),
        ('Registro Activo', lambda o: 'Sí' if o.record_active else 'No'),
    ]


class InvoiceCSVExportView(CSVExportView):
    model = Invoice
    permission_required = 'dashboard.view_invoice'
    filename = 'facturas.csv'
    columns = [
        ('Número', 'number'),
        ('Fecha', lambda o: o.issue_date.isoformat() if o.issue_date else ''),
        ('Cliente', lambda o: o.subscription.customer.company_name if o.subscription and o.subscription.customer else o.customer.company_name if o.customer else ''),
        ('Monto', 'amount'),
        ('Anulada', lambda o: 'Sí' if o.is_cancelled else 'No'),
        ('Correo Enviado', lambda o: 'Sí' if o.email_sent else 'No'),
    ]


class ForecastsCSVExportView(CSVExportView):
    model = Forecasts
    permission_required = 'dashboard.view_forecast'
    filename = 'pronosticos.csv'

    def get_queryset(self):
        qs = self.model.objects.prefetch_related('regions', 'extended_days')
        date_filter = self.request.GET.get('date')
        if date_filter:
            try:
                parsed = datetime.datetime.strptime(date_filter, '%Y-%m-%d').date()
                qs = qs.filter(date=parsed)
                self.filename = f'pronosticos_{date_filter}.csv'
            except (ValueError, TypeError):
                pass
        return qs

    def get_columns(self):
        period_keys = ['morning', 'afternoon', 'night']
        period_labels = ['Mañana', 'Tarde', 'Noche']
        zone_config = [
            ('Costa Norte', 'north', True),
            ('Interior', 'interior', False),
            ('Costa Sur', 'south', True),
        ]

        columns = [
            ('Fecha', lambda o: o.date.isoformat() if o.date else ''),
        ]

        for zone_label, zone_key, has_sea in zone_config:
            for pk, pl in zip(period_keys, period_labels):
                columns += [
                    (f'{zone_label} {pl} Temp',
                     lambda o, zk=zone_key, pk=pk: next(
                         (str(r.temp) for r in o.regions.all() if r.region == zk and r.period == pk), '')),
                    (f'{zone_label} {pl} Tiempo',
                     lambda o, zk=zone_key, pk=pk: next(
                         (r.weather for r in o.regions.all() if r.region == zk and r.period == pk), '')),
                    (f'{zone_label} {pl} Viento DD',
                     lambda o, zk=zone_key, pk=pk: next(
                         (r.wind_dir for r in o.regions.all() if r.region == zk and r.period == pk), '')),
                    (f'{zone_label} {pl} Viento FF',
                     lambda o, zk=zone_key, pk=pk: next(
                         (r.wind_speed for r in o.regions.all() if r.region == zk and r.period == pk), '')),
                ]
                if has_sea:
                    columns.append(
                        (f'{zone_label} {pl} Mar',
                         lambda o, zk=zone_key, pk=pk: next(
                             (r.sea_note or '' for r in o.regions.all() if r.region == zk and r.period == pk), '')))

        for day_num in range(1, 6):
            columns += [
                (f'Ext Día {day_num} Fecha',
                 lambda o, dn=day_num: next(
                     (d.date.isoformat() for d in o.extended_days.all() if d.day_number == dn), '')),
                (f'Ext Día {day_num} Mín',
                 lambda o, dn=day_num: next(
                     (str(d.min_temp) for d in o.extended_days.all() if d.day_number == dn), '')),
                (f'Ext Día {day_num} Máx',
                 lambda o, dn=day_num: next(
                     (str(d.max_temp) for d in o.extended_days.all() if d.day_number == dn), '')),
                (f'Ext Día {day_num} Tiempo',
                 lambda o, dn=day_num: next(
                     (d.weather for d in o.extended_days.all() if d.day_number == dn), '')),
            ]

        columns += [
            ('Fase Lunar', lambda o: o.get_lp_display() if o.lp else ''),
            ('Próxima Fase', lambda o: o.get_nlp_display() if o.nlp else ''),
            ('Fecha Próx Fase', lambda o: o.nlpd.isoformat() if o.nlpd else ''),
            ('Salida Sol', lambda o: o.sunrise.strftime('%H:%M') if o.sunrise else ''),
            ('Puesta Sol', lambda o: o.sunset.strftime('%H:%M') if o.sunset else ''),
            ('Índice UV', lambda o: str(o.uv_index) if o.uv_index is not None else ''),
        ]
        return columns


class WeatherReportCSVExportView(CSVExportView):
    model = WeatherReport
    permission_required = 'dashboard.view_forecast'
    filename = 'reportes_tiempo.csv'
    columns = [
        ('Tipo', lambda o: o.get_type_display()),
        ('Fecha', lambda o: o.date.isoformat() if o.date else ''),
        ('Hora', lambda o: o.time.isoformat() if hasattr(o, 'time') and o.time else ''),
        ('Autor', lambda o: o.author.username if o.author else ''),
        ('Creado', lambda o: o.created_at.isoformat() if hasattr(o, 'created_at') and o.created_at else ''),
    ]


class EarlyWarningCSVExportView(CSVExportView):
    model = EarlyWarning
    permission_required = 'dashboard.view_early_warning'
    filename = 'alertas_tempranas.csv'
    columns = [
        ('Resumen', 'summary'),
        ('Fecha', lambda o: o.date.isoformat() if o.date else ''),
        ('Válido Hasta', lambda o: o.valid_until.isoformat() if hasattr(o, 'valid_until') and o.valid_until else ''),
        ('Usuario', lambda o: o.user.username if o.user else ''),
    ]


class TropicalCycloneCSVExportView(CSVExportView):
    model = TropicalCyclone
    permission_required = 'dashboard.view_tropical_cyclone'
    filename = 'ciclones_tropicales.csv'
    columns = [
        ('Resumen', 'summary'),
        ('Fecha', lambda o: o.date.isoformat() if o.date else ''),
        ('Válido Hasta', lambda o: o.valid_until.isoformat() if hasattr(o, 'valid_until') and o.valid_until else ''),
        ('Usuario', lambda o: o.user.username if o.user else ''),
    ]


class StormWarningCSVExportView(CSVExportView):
    model = StormWarning
    permission_required = 'dashboard.view_storm_warning'
    filename = 'avisos_tormentas.csv'
    columns = [
        ('Resumen', 'summary'),
        ('Fecha', lambda o: o.date.isoformat() if o.date else ''),
        ('Válido Hasta', lambda o: o.valid_until.isoformat() if hasattr(o, 'valid_until') and o.valid_until else ''),
        ('Usuario', lambda o: o.user.username if o.user else ''),
    ]


class ForecastsExcelExportView(LoginRequiredMixin, PermissionRequiredMixin, View):
    model = Forecasts
    permission_required = 'dashboard.view_forecast'

    def get_forecasts(self):
        qs = Forecasts.objects.prefetch_related('regions', 'extended_days')
        date_filter = self.request.GET.get('date')
        if date_filter:
            try:
                parsed = datetime.datetime.strptime(date_filter, '%Y-%m-%d').date()
                qs = qs.filter(date=parsed)
            except (ValueError, TypeError):
                pass
        return qs.order_by('date')

    def _style_cell(self, cell, bold=False):
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin'),
        )
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center', vertical='center')
        if bold:
            cell.font = Font(bold=True)

    def _write_forecast_sheet(self, ws, forecast):
        zone_map = {'north': 'Costa Norte', 'interior': 'Interior', 'south': 'Costa Sur'}
        period_keys = ['morning', 'afternoon', 'night']
        period_labels = ['Mañana', 'Tarde', 'Noche']
        col_groups = [(2, 'Temperatura'), (5, 'Tiempo'), (8, 'Viento (DD)'), (11, 'Viento (FF)'), (14, 'Mar')]

        # Fecha
        ws['A1'] = 'Fecha'
        self._style_cell(ws['A1'], bold=True)
        ws['B1'] = forecast.date.isoformat() if forecast.date else ''
        self._style_cell(ws['B1'])

        # Pronóstico label
        ws['A3'] = 'Pronóstico'
        self._style_cell(ws['A3'], bold=True)

        # Headers row 4
        ws['A4'] = 'Zona'
        self._style_cell(ws['A4'], bold=True)
        for col_start, header in col_groups:
            cell = ws.cell(row=4, column=col_start)
            cell.value = header
            self._style_cell(cell, bold=True)
            ws.merge_cells(start_row=4, start_column=col_start, end_row=4, end_column=col_start + 2)

        ws.merge_cells('A4:A5')

        # Sub-headers row 5
        for col_start, _ in col_groups:
            for i, label in enumerate(period_labels):
                cell = ws.cell(row=5, column=col_start + i)
                cell.value = label
                self._style_cell(cell, bold=True)

        # Region data rows 6-8
        regions_by_zone = {}
        for reg in forecast.regions.all():
            regions_by_zone.setdefault(reg.region, {})[reg.period] = reg

        for row_offset, (zone_key, has_sea) in enumerate([('north', True), ('interior', False), ('south', True)]):
            row = 6 + row_offset
            ws.cell(row=row, column=1).value = zone_map[zone_key]
            self._style_cell(ws.cell(row=row, column=1))

            zone_data = regions_by_zone.get(zone_key, {})
            field_map = [
                (2, 'temp'), (5, 'weather'), (8, 'wind_dir'), (11, 'wind_speed'), (14, 'sea_note'),
            ]
            for col_start, field in field_map:
                for i, pk in enumerate(period_keys):
                    reg = zone_data.get(pk)
                    val = getattr(reg, field, '') if reg else ''
                    cell = ws.cell(row=row, column=col_start + i)
                    cell.value = val if val is not None else ''
                    self._style_cell(cell)

        # Extended forecast header row 11
        ws['A10'] = 'Pronóstico Extendido'
        self._style_cell(ws['A10'], bold=True)
        ws['F10'] = 'Datos Astronómicos'
        self._style_cell(ws['F10'], bold=True)

        ext_headers = ['Día', 'Minima', 'Máxima', 'Tiempo']
        for col_num, h in enumerate(ext_headers, start=1):
            cell = ws.cell(row=11, column=col_num)
            cell.value = h
            self._style_cell(cell, bold=True)

        # Extended forecast data rows 12-16
        extended_by_day = {d.day_number: d for d in forecast.extended_days.all()}
        for day_num in range(1, 6):
            row = 11 + day_num
            ed = extended_by_day.get(day_num)
            vals = [
                ed.date.isoformat() if ed and ed.date else '',
                str(ed.min_temp) if ed and ed.min_temp is not None else '',
                str(ed.max_temp) if ed and ed.max_temp is not None else '',
                ed.weather if ed else '',
            ]
            for col_num, v in enumerate(vals, start=1):
                cell = ws.cell(row=row, column=col_num)
                cell.value = v
                self._style_cell(cell)

        # Astro headers row 11 (cols 7-9)
        astro_headers = ['Actual', 'Próxima', 'Fecha']
        for col_num, h in enumerate(astro_headers, start=7):
            cell = ws.cell(row=11, column=col_num)
            cell.value = h
            self._style_cell(cell, bold=True)

        # Astro data rows 12-15 (cols 6-9)
        astro_data = [
            ['Fase Lunar', forecast.get_lp_display() if forecast.lp else '',
             forecast.get_nlp_display() if forecast.nlp else '',
             forecast.nlpd.isoformat() if forecast.nlpd else ''],
            ['Salida Sol', forecast.sunrise.strftime('%H:%M') if forecast.sunrise else '', '', ''],
            ['Puesta Sol', forecast.sunset.strftime('%H:%M') if forecast.sunset else '', '', ''],
            ['Índice UV', str(forecast.uv_index) if forecast.uv_index is not None else '', '', ''],
        ]
        for i, row_data in enumerate(astro_data):
            row = 12 + i
            for j, val in enumerate(row_data, start=6):
                cell = ws.cell(row=row, column=j)
                cell.value = val
                self._style_cell(cell)

        # Auto-adjust column widths
        for col_cells in ws.columns:
            max_len = 0
            col_letter = col_cells[0].column_letter
            for cell in col_cells:
                try:
                    if cell.value and len(str(cell.value)) > max_len:
                        max_len = len(str(cell.value))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = max_len + 3

    def get(self, request):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        forecasts = self.get_forecasts()
        for forecast in forecasts:
            date_str = forecast.date.isoformat() if forecast.date else 'sin_fecha'
            ws = wb.create_sheet(title=date_str[:31])
            self._write_forecast_sheet(ws, forecast)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        date_param = self.request.GET.get('date', 'todos')
        response['Content-Disposition'] = f'attachment; filename="pronosticos_{date_param}.xlsx"'
        wb.save(response)
        return response
